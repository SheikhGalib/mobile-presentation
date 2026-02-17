from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Colors
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
DARK_GREEN = RGBColor(0, 100, 0)
GREEN = RGBColor(34, 139, 34)
LIGHT_GREEN = RGBColor(144, 238, 144)
ACCENT_GREEN = RGBColor(39, 174, 96)
DARK_GRAY = RGBColor(60, 60, 60)
LIGHT_GRAY = RGBColor(200, 200, 200)
VERY_LIGHT_GRAY = RGBColor(245, 245, 245)
MID_GRAY = RGBColor(120, 120, 120)
RED_ACCENT = RGBColor(192, 57, 43)
BLUE_ACCENT = RGBColor(41, 128, 185)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_green_bar(slide, left=0, top=0, width=None, height=Inches(0.08)):
    """Add a thin green accent bar"""
    if width is None:
        width = prs.slide_width
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_GREEN
    shape.line.fill.background()
    return shape

def add_bottom_bar(slide):
    """Add bottom green bar with page number area"""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.1), prs.slide_width, Inches(0.4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK_GREEN
    bar.line.fill.background()
    return bar

def add_side_accent(slide, left=0, top=Inches(1.2), width=Inches(0.06), height=Inches(5.5)):
    """Add a thin green side accent"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_GREEN
    shape.line.fill.background()
    return shape

def set_slide_bg_white(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE

def add_text_box(slide, left, top, width, height, text, font_size=18, bold=False, color=BLACK, alignment=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_bullet_points(slide, left, top, width, height, items, font_size=16, color=BLACK, spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = 'Calibri'
        p.space_after = spacing
        p.level = 0
    return txBox

def add_icon_box(slide, left, top, size, text, bg_color=DARK_GREEN, text_color=WHITE, font_size=14):
    """Add a colored box with icon/text"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.name = 'Calibri'
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return shape

def add_rounded_box(slide, left, top, width, height, text, bg_color=VERY_LIGHT_GRAY, border_color=DARK_GREEN, font_size=14, text_color=BLACK, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.name = 'Calibri'
    p.font.bold = bold
    return shape

def add_arrow_right(slide, left, top, width=Inches(0.5), height=Inches(0.3)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_GREEN
    shape.line.fill.background()
    return shape

def add_slide_number(slide, num, total=12):
    txt = add_text_box(slide, Inches(12.3), Inches(7.15), Inches(0.9), Inches(0.3),
                       f"{num}/{total}", font_size=11, color=WHITE, alignment=PP_ALIGN.RIGHT)

# ============================================================
# SLIDE 1: TITLE SLIDE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg_white(slide)

# Top green bar
add_green_bar(slide, top=0, height=Inches(0.12))

# Bottom green bar (thicker for title slide)
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.5), prs.slide_width, Inches(1.0))
bar.fill.solid()
bar.fill.fore_color.rgb = DARK_GREEN
bar.line.fill.background()

# Left accent block
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.5), Inches(0.5), Inches(3.0))
accent.fill.solid()
accent.fill.fore_color.rgb = ACCENT_GREEN
accent.line.fill.background()

# Title
add_text_box(slide, Inches(1.0), Inches(1.5), Inches(10), Inches(1.2),
             "Designing for a Unified Experience", font_size=40, bold=True, color=DARK_GREEN, font_name='Calibri')
# Subtitle
add_text_box(slide, Inches(1.0), Inches(2.6), Inches(10), Inches(0.6),
             "A New Perspective and a Case Study", font_size=24, color=DARK_GRAY, font_name='Calibri')

# Divider line
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.3), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# Paper authors
add_text_box(slide, Inches(1.0), Inches(3.5), Inches(10), Inches(0.4),
             "Paper by: Wei Xu & Dov Furie — Intel Corporation", font_size=14, color=MID_GRAY)

# Course info
add_text_box(slide, Inches(1.0), Inches(4.2), Inches(5), Inches(0.4),
             "Course: CSE [Course No.] — [Course Name]", font_size=16, color=DARK_GRAY, bold=True)
add_text_box(slide, Inches(1.0), Inches(4.6), Inches(5), Inches(0.4),
             "Course Teacher: [Teacher Name]", font_size=14, color=DARK_GRAY)

# Student Info
add_text_box(slide, Inches(7.0), Inches(4.2), Inches(5), Inches(0.3),
             "Presented by:", font_size=14, bold=True, color=DARK_GREEN)
students = [
    "1. [Student Name 1] — Roll: [XXXXXXX]",
    "2. [Student Name 2] — Roll: [XXXXXXX]",
    "3. [Student Name 3] — Roll: [XXXXXXX]",
    "4. [Student Name 4] — Roll: [XXXXXXX]",
]
add_bullet_points(slide, Inches(7.0), Inches(4.5), Inches(5), Inches(1.5), students, font_size=13, color=DARK_GRAY)

# KUET logo placeholder
logo_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.5), Inches(1.0), Inches(1.3), Inches(1.3))
logo_box.fill.solid()
logo_box.fill.fore_color.rgb = VERY_LIGHT_GRAY
logo_box.line.color.rgb = DARK_GREEN
logo_box.line.width = Pt(2)
tf = logo_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "KUET\nLogo"
p.font.size = Pt(12)
p.font.color.rgb = DARK_GREEN
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

# Bottom bar text
add_text_box(slide, Inches(0.5), Inches(6.6), Inches(12), Inches(0.6),
             "Khulna University of Engineering & Technology    |    Department of Computer Science & Engineering    |    February 2026",
             font_size=13, color=WHITE, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 2: TABLE OF CONTENTS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 2)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "Roadmap", font_size=32, bold=True, color=DARK_GREEN)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(2.5), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# Roadmap items as styled boxes
toc_items = [
    ("01", "The Problem", "Why Enterprise UX is Broken"),
    ("02", "Previous Efforts", "What Others Tried (& Why It Wasn't Enough)"),
    ("03", "Unified Experience", "The 5-Attribute Framework"),
    ("04", "Implementation Framework", "2 Driving Vectors"),
    ("05", "Case Study", "Meet Paul — The Frustrated Buyer"),
    ("06", "Results", "Hard Numbers That Prove It Works"),
    ("07", "Takeaways", "Key Lessons & Exam Points"),
]

y_start = Inches(1.5)
for i, (num, title, desc) in enumerate(toc_items):
    y = y_start + Inches(i * 0.75)
    
    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), y, Inches(0.55), Inches(0.55))
    circle.fill.solid()
    circle.fill.fore_color.rgb = DARK_GREEN
    circle.line.fill.background()
    tf = circle.text_frame
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Title
    add_text_box(slide, Inches(1.8), y, Inches(3.5), Inches(0.35),
                 title, font_size=18, bold=True, color=BLACK)
    # Description
    add_text_box(slide, Inches(1.8), y + Inches(0.3), Inches(5), Inches(0.3),
                 desc, font_size=13, color=MID_GRAY)

# Right side decorative element
accent_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.5), Inches(1.5), Inches(4.3), Inches(5.3))
accent_box.fill.solid()
accent_box.fill.fore_color.rgb = VERY_LIGHT_GRAY
accent_box.line.fill.background()

add_text_box(slide, Inches(8.8), Inches(1.8), Inches(3.8), Inches(1.0),
             "5-Minute Presentation", font_size=22, bold=True, color=DARK_GREEN, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(8.8), Inches(2.6), Inches(3.8), Inches(3.5),
             "This paper proposes a Unified Experience framework for enterprise software, backed by a real Intel case study that dramatically improved worker productivity.",
             font_size=14, color=DARK_GRAY, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(8.8), Inches(5.5), Inches(3.8), Inches(0.5),
             "Paper: Xu & Furie, Intel Corp.", font_size=12, color=MID_GRAY, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 3: THE PROBLEM
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 3)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "The Problem — Broken Enterprise UX", font_size=32, bold=True, color=DARK_GREEN)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# Hook question
hook_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.7))
hook_box.fill.solid()
hook_box.fill.fore_color.rgb = RGBColor(255, 248, 230)
hook_box.line.color.rgb = RGBColor(255, 193, 7)
hook_box.line.width = Pt(1.5)
tf = hook_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = '💡 "Imagine using 6+ different apps just to complete ONE task at work... every single day."'
p.font.size = Pt(16)
p.font.color.rgb = DARK_GRAY
p.font.name = 'Calibri'
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

# Left column - The Reality
add_text_box(slide, Inches(0.8), Inches(2.3), Inches(5.5), Inches(0.4),
             "The Reality in Large Companies", font_size=20, bold=True, color=BLACK)

reality_items = [
    "▶ Hundreds of apps: CRM, SCM, ERP, HR, BI...",
    "▶ Hybrid environment: vendor + home-grown apps",
    "▶ Result of mergers, acquisitions over years",
    "▶ Siloed solutions — apps don't talk to each other",
    "▶ Inconsistent UIs, fragmented workflows",
    "▶ Poor data integration across systems",
]
add_bullet_points(slide, Inches(0.8), Inches(2.8), Inches(5.5), Inches(3.0), reality_items, font_size=15, color=DARK_GRAY)

# Right column - Impact
impact_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(2.3), Inches(5.5), Inches(4.5))
impact_box.fill.solid()
impact_box.fill.fore_color.rgb = RGBColor(255, 240, 240)
impact_box.line.color.rgb = RED_ACCENT
impact_box.line.width = Pt(1.5)

add_text_box(slide, Inches(7.3), Inches(2.5), Inches(5.0), Inches(0.4),
             "⚠ Impact on Users", font_size=20, bold=True, color=RED_ACCENT)

impact_items = [
    "✗ Low productivity — time wasted switching apps",
    "✗ High frustration & poor satisfaction",
    "✗ Increased support call volume",
    "✗ Poor decision-making efficiency",
    "✗ High error rates & costly rework",
    "✗ Long learning curves for new employees",
    "✗ Business objectives negatively impacted",
]
add_bullet_points(slide, Inches(7.3), Inches(3.1), Inches(5.0), Inches(3.5), impact_items, font_size=14, color=DARK_GRAY)

# Diagram: Multiple apps → confused user
# Show simplified version
diagram_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.5), Inches(5.0), Inches(1.3))
diagram_box.fill.solid()
diagram_box.fill.fore_color.rgb = VERY_LIGHT_GRAY
diagram_box.line.color.rgb = LIGHT_GRAY

# Mini app boxes
app_names = ["CRM", "ERP", "SCM", "BI", "HR", "Mail"]
for j, name in enumerate(app_names):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 
                                  Inches(1.2 + j*0.7), Inches(5.65), Inches(0.6), Inches(0.45))
    box.fill.solid()
    box.fill.fore_color.rgb = [RGBColor(231, 76, 60), RGBColor(52, 152, 219), RGBColor(46, 204, 113), 
                                RGBColor(155, 89, 182), RGBColor(241, 196, 15), RGBColor(230, 126, 34)][j]
    box.line.fill.background()
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = name
    p.font.size = Pt(9)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

add_text_box(slide, Inches(1.0), Inches(6.15), Inches(5.0), Inches(0.5),
             "User must juggle ALL of these daily → Broken Experience", font_size=12, color=RED_ACCENT, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 4: PREVIOUS APPROACHES
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 4)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Previous Approaches — What Others Tried", font_size=32, bold=True, color=DARK_GREEN)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# Table header
header_y = Inches(1.4)
cols = [Inches(0.8), Inches(3.5), Inches(6.5), Inches(9.8)]
col_widths = [Inches(2.5), Inches(2.8), Inches(3.1), Inches(3.0)]
headers = ["Approach", "Who", "What They Did", "Limitation"]

for i, (col, w, h) in enumerate(zip(cols, col_widths, headers)):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col, header_y, w, Inches(0.4))
    box.fill.solid()
    box.fill.fore_color.rgb = DARK_GREEN
    box.line.fill.background()
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = h
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

# Table rows
rows_data = [
    ["Uniform Visual\nDesign", "Google, Cooper", "Consistent look &\nfeel across apps", "Surface-level only —\ndoesn't fix workflows"],
    ["Unified Customer\nExperience", "Rogowski /\nForrester", "Seamless cross-channel\ntouchpoints", "Limited to digital\nmarketing domain"],
    ["Unified\nArchitecture", "O'Toole", "Integrated backend\n& database", "No front-end or\nprocess details"],
    ["Vendor\nPartnership", "Finstad, Xu", "Bridge vendor-user\ngap", "Not full end-to-end\nUX coverage"],
    ["Klocek's Hierarchy\nModel", "Klocek", "5-level pyramid:\nvisual → org shift", "Too linear, visual\noveremphasized"],
]

for r, row in enumerate(rows_data):
    row_y = header_y + Inches(0.45) + Inches(r * 0.65)
    bg = VERY_LIGHT_GRAY if r % 2 == 0 else WHITE
    for i, (col, w, cell) in enumerate(zip(cols, col_widths, row)):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col, row_y, w, Inches(0.6))
        box.fill.solid()
        box.fill.fore_color.rgb = bg
        box.line.color.rgb = LIGHT_GRAY
        box.line.width = Pt(0.5)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = cell
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p.font.name = 'Calibri'
        p.alignment = PP_ALIGN.CENTER

# Klocek criticism box
crit_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.7))
crit_box.fill.solid()
crit_box.fill.fore_color.rgb = RGBColor(255, 248, 240)
crit_box.line.color.rgb = RGBColor(230, 126, 34)
crit_box.line.width = Pt(1.5)

add_text_box(slide, Inches(1.1), Inches(5.3), Inches(5), Inches(0.4),
             "Key Criticism of Klocek's 5-Level Model:", font_size=16, bold=True, color=RGBColor(230, 126, 34))

crit_items = [
    "1. Overemphasizes visual design as the foundation — but pretty UI alone doesn't fix broken processes",
    "2. Doesn't fit hybrid vendor environments — you often CAN'T customize vendor apps",
    "3. Business process redesign is placed too late — should be earlier priority",
    "4. Too linear — real-world implementation doesn't follow a neat bottom-up or top-down sequence",
]
add_bullet_points(slide, Inches(1.1), Inches(5.8), Inches(11.0), Inches(1.2), crit_items, font_size=12, color=DARK_GRAY, spacing=Pt(3))


# ============================================================
# SLIDE 5: UNIFIED EXPERIENCE - 5 ATTRIBUTES (CORE!)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 5)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "The Solution — 5 Attributes of Unified Experience", font_size=32, bold=True, color=DARK_GREEN)

star = add_text_box(slide, Inches(11.5), Inches(0.3), Inches(1.5), Inches(0.7),
             "★ CORE", font_size=18, bold=True, color=RED_ACCENT, alignment=PP_ALIGN.CENTER)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# Center hexagon-like arrangement
attrs = [
    ("1", "Consistency", "Uniform visual &\nbranding design\nacross applications", Inches(5.5), Inches(1.4)),
    ("2", "Efficiency", "Streamlined\nbusiness processes\nfor quality & savings", Inches(2.0), Inches(2.8)),
    ("3", "Personalization", "Right content,\nright device,\nright context", Inches(9.0), Inches(2.8)),
    ("4", "Integration", "Data, security &\narchitecture\nacross apps", Inches(2.8), Inches(5.0)),
    ("5", "Optimization", "Seamless end-to-end\ntouchpoints:\nperformance, help, support", Inches(8.2), Inches(5.0)),
]

for num, title, desc, x, y in attrs:
    # Main box
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.2), Inches(1.5))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = DARK_GREEN
    box.line.width = Pt(2)
    
    # Number badge
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - Inches(0.15), y - Inches(0.15), Inches(0.45), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = DARK_GREEN
    badge.line.fill.background()
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Title
    add_text_box(slide, x + Inches(0.2), y + Inches(0.1), Inches(2.8), Inches(0.35),
                 title, font_size=17, bold=True, color=DARK_GREEN, alignment=PP_ALIGN.CENTER)
    # Description
    add_text_box(slide, x + Inches(0.2), y + Inches(0.5), Inches(2.8), Inches(0.9),
                 desc, font_size=12, color=DARK_GRAY, alignment=PP_ALIGN.CENTER)

# Central "UNIFIED EXPERIENCE" label
center_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(3.5), Inches(3.7), Inches(1.0))
center_box.fill.solid()
center_box.fill.fore_color.rgb = DARK_GREEN
center_box.line.fill.background()
tf = center_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "UNIFIED\nEXPERIENCE"
p.font.size = Pt(18)
p.font.color.rgb = WHITE
p.font.bold = True
p.alignment = PP_ALIGN.CENTER


# ============================================================
# SLIDE 6: IMPLEMENTATION FRAMEWORK
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 6)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Implementation Framework — 2 Driving Vectors", font_size=32, bold=True, color=DARK_GREEN)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# TOP-DOWN arrow
top_arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(0.8), Inches(1.5), Inches(1.0), Inches(2.5))
top_arrow.fill.solid()
top_arrow.fill.fore_color.rgb = RGBColor(41, 128, 185)
top_arrow.line.fill.background()
tf = top_arrow.text_frame
p = tf.paragraphs[0]
p.text = "TOP\nDOWN"
p.font.size = Pt(11)
p.font.color.rgb = WHITE
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

# Organization boxes
add_text_box(slide, Inches(2.0), Inches(1.5), Inches(4), Inches(0.4),
             "Vector 1: Organization (Top-Down)", font_size=18, bold=True, color=BLUE_ACCENT)

org_items = [
    ("UX Culture", "Experience scorecard,\nUX maturity model"),
    ("Organizational\nCommitment", "Drive UX-based investment,\nstrategic procurement"),
]
for i, (title, desc) in enumerate(org_items):
    y = Inches(2.0) + Inches(i * 1.0)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), y, Inches(2.0), Inches(0.8))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(214, 234, 248)
    box.line.color.rgb = BLUE_ACCENT
    box.line.width = Pt(1.5)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p.alignment = PP_ALIGN.CENTER
    
    add_text_box(slide, Inches(4.2), y, Inches(2.5), Inches(0.8), desc, font_size=11, color=DARK_GRAY)

# BOTTOM-UP arrow
bottom_arrow = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(0.8), Inches(4.2), Inches(1.0), Inches(2.7))
bottom_arrow.fill.solid()
bottom_arrow.fill.fore_color.rgb = ACCENT_GREEN
bottom_arrow.line.fill.background()
tf = bottom_arrow.text_frame
p = tf.paragraphs[0]
p.text = "BOTTOM\nUP"
p.font.size = Pt(11)
p.font.color.rgb = WHITE
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

# Design boxes
add_text_box(slide, Inches(2.0), Inches(4.2), Inches(4), Inches(0.4),
             "Vector 2: Design (Bottom-Up)", font_size=18, bold=True, color=DARK_GREEN)

design_items = [
    ("UX Standards", "Design standards &\nguidelines"),
    ("UI Assets", "Reusable style sheets,\nUI components"),
    ("Integration Framework", "API services,\ntoolkits"),
    ("UX Practices", "Agile development\nmethods & tools"),
    ("Vendor Strategy", "Consolidation,\nvendor assessment"),
    ("Governance", "Checklists,\nautomatic tools"),
]

for i, (title, desc) in enumerate(design_items):
    col = i % 3
    row = i // 3
    x = Inches(2.0) + Inches(col * 2.2)
    y = Inches(4.7) + Inches(row * 1.0)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.0), Inches(0.8))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(234, 250, 234)
    box.line.color.rgb = ACCENT_GREEN
    box.line.width = Pt(1.5)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.alignment = PP_ALIGN.CENTER

# Right side: Key features
feature_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.5), Inches(4.3), Inches(5.3))
feature_box.fill.solid()
feature_box.fill.fore_color.rgb = VERY_LIGHT_GRAY
feature_box.line.color.rgb = DARK_GREEN
feature_box.line.width = Pt(1.5)

add_text_box(slide, Inches(8.8), Inches(1.7), Inches(3.8), Inches(0.4),
             "Why This Framework is Better:", font_size=16, bold=True, color=DARK_GREEN)

features = [
    "✓  End-to-end systematic view",
    "✓  Holistic — all user touchpoints",
    "✓  Realistic — works with vendor solutions",
    "✓  Adaptive — start from anywhere",
    "✓  UX metric-driven — prove ROI",
    "✓  Tactical — specific tools & approaches",
]
add_bullet_points(slide, Inches(8.8), Inches(2.2), Inches(3.8), Inches(3.5), features, font_size=14, color=DARK_GRAY, spacing=Pt(10))


# ============================================================
# SLIDE 7: CASE STUDY - MEET PAUL
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 7)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Case Study — Meet Paul, the Frustrated Buyer", font_size=32, bold=True, color=DARK_GREEN)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# Paul persona box
persona_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(4.0), Inches(5.3))
persona_box.fill.solid()
persona_box.fill.fore_color.rgb = RGBColor(240, 248, 255)
persona_box.line.color.rgb = BLUE_ACCENT
persona_box.line.width = Pt(2)

# Person icon
person = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.3), Inches(1.6), Inches(0.8), Inches(0.8))
person.fill.solid()
person.fill.fore_color.rgb = BLUE_ACCENT
person.line.fill.background()
tf = person.text_frame
p = tf.paragraphs[0]
p.text = "👤"
p.font.size = Pt(20)
p.alignment = PP_ALIGN.CENTER

add_text_box(slide, Inches(1.0), Inches(2.5), Inches(3.6), Inches(0.4),
             "Paul — Supply Chain Buyer", font_size=16, bold=True, color=BLUE_ACCENT, alignment=PP_ALIGN.CENTER)

paul_info = [
    "▸ Works at global company",
    "▸ 10 years experience, 3.5 in current role",
    "▸ Manages factory spare parts inventory",
    "▸ Day driven by tactics & exceptions",
    "▸ Constantly fighting fires",
    "▸ Liaison between factory & supplier",
    "▸ Wishes: 'ONE system for all tasks'",
]
add_bullet_points(slide, Inches(1.0), Inches(3.0), Inches(3.6), Inches(3.0), paul_info, font_size=13, color=DARK_GRAY, spacing=Pt(5))

# Pain points - right side
add_text_box(slide, Inches(5.2), Inches(1.4), Inches(7.5), Inches(0.4),
             "Paul's Daily Nightmare:", font_size=20, bold=True, color=RED_ACCENT)

# Stat boxes
stats = [
    ("6+", "Different\nApplications"),
    ("35", "Manual\nSteps"),
    ("69", "Mouse Clicks\nper Task"),
    ("10 min", "Average\nTask Time"),
]

for i, (num, label) in enumerate(stats):
    x = Inches(5.2) + Inches(i * 2.0)
    # Number
    num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(1.7), Inches(0.7))
    num_box.fill.solid()
    num_box.fill.fore_color.rgb = RED_ACCENT
    num_box.line.fill.background()
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(24)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    add_text_box(slide, x, Inches(2.8), Inches(1.7), Inches(0.6),
                 label, font_size=11, color=DARK_GRAY, alignment=PP_ALIGN.CENTER)

# Problems table
add_text_box(slide, Inches(5.2), Inches(3.7), Inches(7.5), Inches(0.4),
             "Three Core Problems:", font_size=16, bold=True, color=BLACK)

prob_headers = ["Issue", "Description", "Severity"]
prob_cols = [Inches(5.2), Inches(7.2), Inches(10.5)]
prob_widths = [Inches(1.8), Inches(3.1), Inches(1.5)]

for i, (col, w, h) in enumerate(zip(prob_cols, prob_widths, prob_headers)):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col, Inches(4.2), w, Inches(0.35))
    box.fill.solid()
    box.fill.fore_color.rgb = DARK_GREEN
    box.line.fill.background()
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = h
    p.font.size = Pt(11)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

problems = [
    ["Long Cycle Time", "Long waiting + processing time → low productivity", "HIGH"],
    ["Wrong Priorities", "Critical items not worked first → factory escalations", "HIGH"],
    ["Quality Issues", "Incorrect transactions → rework and cost", "HIGH"],
]

for r, row in enumerate(problems):
    y = Inches(4.6) + Inches(r * 0.55)
    bg = VERY_LIGHT_GRAY if r % 2 == 0 else WHITE
    for i, (col, w, cell) in enumerate(zip(prob_cols, prob_widths, row)):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col, y, w, Inches(0.5))
        box.fill.solid()
        box.fill.fore_color.rgb = bg
        box.line.color.rgb = LIGHT_GRAY
        box.line.width = Pt(0.5)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = cell
        p.font.size = Pt(11)
        fc = RED_ACCENT if i == 2 else DARK_GRAY
        p.font.color.rgb = fc
        p.font.bold = (i == 2)
        p.alignment = PP_ALIGN.CENTER

# Quote
add_text_box(slide, Inches(5.2), Inches(6.3), Inches(7.0), Inches(0.5),
             '"I wish to have ONE comprehensive system for all transactions instead of dozen different systems."',
             font_size=12, color=MID_GRAY, alignment=PP_ALIGN.LEFT)


# ============================================================
# SLIDE 8: DESIGN PRINCIPLES
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 8)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "6 Design Principles — How They Fixed It", font_size=32, bold=True, color=DARK_GREEN)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

principles = [
    ("1", "Standardize\nthe Process", "Create a single version\nof the process", "UX, Productivity"),
    ("2", "Integrate\nBusiness Intelligence", "Minimum info for\ndecision on ONE screen", "Decision Making"),
    ("3", "Reduce\nSiloed Systems", "Integrate data from\nmultiple sources", "Focus, Accuracy"),
    ("4", "Collaborate\nOnline", "Replace email with\ndigital collaboration", "Speed, Automation"),
    ("5", "Eliminate\nRework", "Verify data at\npoint of entry", "Quality, Productivity"),
    ("6", "Self-Service\nfor Customers", "Expose process state\nto stakeholders", "Trust, Efficiency"),
]

for i, (num, title, desc, benefit) in enumerate(principles):
    col = i % 3
    row = i // 3
    x = Inches(0.8) + Inches(col * 4.1)
    y = Inches(1.5) + Inches(row * 2.3)
    
    # Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.8), Inches(2.0))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = ACCENT_GREEN
    card.line.width = Pt(1.5)
    
    # Number badge
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), y + Inches(0.15), Inches(0.4), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = DARK_GREEN
    badge.line.fill.background()
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Title
    add_text_box(slide, x + Inches(0.7), y + Inches(0.1), Inches(2.8), Inches(0.55),
                 title, font_size=14, bold=True, color=DARK_GREEN)
    # Description
    add_text_box(slide, x + Inches(0.2), y + Inches(0.75), Inches(3.4), Inches(0.6),
                 desc, font_size=12, color=DARK_GRAY)
    # Benefit tag
    tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y + Inches(1.5), Inches(2.0), Inches(0.3))
    tag.fill.solid()
    tag.fill.fore_color.rgb = RGBColor(234, 250, 234)
    tag.line.fill.background()
    tf = tag.text_frame
    p = tf.paragraphs[0]
    p.text = f"→ {benefit}"
    p.font.size = Pt(10)
    p.font.color.rgb = DARK_GREEN
    p.alignment = PP_ALIGN.CENTER

# Tech used box
tech_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.7))
tech_box.fill.solid()
tech_box.fill.fore_color.rgb = RGBColor(240, 248, 255)
tech_box.line.color.rgb = BLUE_ACCENT
tech_box.line.width = Pt(1.5)

add_text_box(slide, Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.5),
             "Technology: BPMS (Business Process Management Suite) + SOA (Service Oriented Architecture) → New system: BOMA",
             font_size=14, bold=True, color=BLUE_ACCENT, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 9: RESULTS - BEFORE vs AFTER
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 9)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "Results — The Numbers That Prove It Works", font_size=32, bold=True, color=DARK_GREEN)

star2 = add_text_box(slide, Inches(10.5), Inches(0.3), Inches(2.5), Inches(0.7),
             "📊 KEY SLIDE", font_size=18, bold=True, color=RED_ACCENT, alignment=PP_ALIGN.CENTER)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# LEFT: Process/Integration/Optimization table
add_text_box(slide, Inches(0.8), Inches(1.3), Inches(5.5), Inches(0.4),
             "UX Improvements (Before → After)", font_size=18, bold=True, color=BLACK)

# Table headers
table_left = Inches(0.8)
col1_w = Inches(1.5)
col2_w = Inches(2.3)
col3_w = Inches(0.8)
col4_w = Inches(0.8)

th_y = Inches(1.8)
for col, w, text in [(table_left, col1_w, "Category"), 
                      (table_left+col1_w, col2_w, "Metric"),
                      (table_left+col1_w+col2_w, col3_w, "Before"),
                      (table_left+col1_w+col2_w+col3_w, col4_w, "After")]:
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col, th_y, w, Inches(0.35))
    box.fill.solid()
    box.fill.fore_color.rgb = DARK_GREEN
    box.line.fill.background()
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

left_rows = [
    ("Process", "Manual Steps", "35", "12"),
    ("", "Manual Processes", "7", "1"),
    ("", "Roles", "5", "4"),
    ("", "Rework Rate", "20%", "<5%"),
    ("Integration", "Applications", "6", "1"),
    ("", "Screens", "15", "3"),
    ("", "Automated Steps", "8", "23"),
    ("Optimization", "Manual Emails", "Dozens", "None"),
    ("", "Online Help", "Partial", "Full"),
]

for r, (cat, metric, before, after) in enumerate(left_rows):
    y = th_y + Inches(0.38) + Inches(r * 0.38)
    bg = VERY_LIGHT_GRAY if r % 2 == 0 else WHITE
    
    for col, w, text, is_after in [
        (table_left, col1_w, cat, False),
        (table_left+col1_w, col2_w, metric, False),
        (table_left+col1_w+col2_w, col3_w, before, False),
        (table_left+col1_w+col2_w+col3_w, col4_w, after, True),
    ]:
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, col, y, w, Inches(0.35))
        box.fill.solid()
        box.fill.fore_color.rgb = bg
        box.line.color.rgb = LIGHT_GRAY
        box.line.width = Pt(0.5)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(10)
        if is_after:
            p.font.color.rgb = DARK_GREEN
            p.font.bold = True
        elif cat:
            p.font.color.rgb = DARK_GREEN
            p.font.bold = True
        else:
            p.font.color.rgb = DARK_GRAY
        p.alignment = PP_ALIGN.CENTER

# RIGHT: Business Value metrics
add_text_box(slide, Inches(7.0), Inches(1.3), Inches(5.5), Inches(0.4),
             "Business Value (3 Months After)", font_size=18, bold=True, color=BLACK)

bv_metrics = [
    ("Task Time", "10 min", "5 min", "50% faster"),
    ("Clicks/Transaction", "69", "17", "75% fewer"),
    ("Late Orders", "40%", "20%", "50% reduction"),
    ("Cost Savings/Buyer", "$0", "$1.2M", "💰"),
    ("Training Time", "Dozens hrs", "2 hours", "95% less"),
    ("'How do I' Tickets", "Dozens", "0", "100% gone"),
    ("System Trust", "<10%", "83%", "8x increase"),
    ("User Feeling", "Frustrated", "Enthusiastic!", "🎉"),
]

for i, (metric, before, after, change) in enumerate(bv_metrics):
    y = Inches(1.8) + Inches(i * 0.62)
    
    # Metric name
    add_text_box(slide, Inches(7.0), y, Inches(2.0), Inches(0.3),
                 metric, font_size=12, bold=True, color=DARK_GRAY)
    
    # Before box
    bb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.1), y, Inches(1.0), Inches(0.4))
    bb.fill.solid()
    bb.fill.fore_color.rgb = RGBColor(255, 235, 235)
    bb.line.fill.background()
    tf = bb.text_frame
    p = tf.paragraphs[0]
    p.text = before
    p.font.size = Pt(10)
    p.font.color.rgb = RED_ACCENT
    p.alignment = PP_ALIGN.CENTER
    
    # Arrow
    add_text_box(slide, Inches(10.15), y, Inches(0.4), Inches(0.4),
                 "→", font_size=14, bold=True, color=DARK_GREEN, alignment=PP_ALIGN.CENTER)
    
    # After box
    ab = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), y, Inches(1.0), Inches(0.4))
    ab.fill.solid()
    ab.fill.fore_color.rgb = RGBColor(234, 250, 234)
    ab.line.fill.background()
    tf = ab.text_frame
    p = tf.paragraphs[0]
    p.text = after
    p.font.size = Pt(10)
    p.font.color.rgb = DARK_GREEN
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Change badge
    add_text_box(slide, Inches(11.6), y, Inches(1.2), Inches(0.4),
                 change, font_size=10, color=ACCENT_GREEN, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 10: KEY TAKEAWAYS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 10)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Key Takeaways & Lessons Learned", font_size=32, bold=True, color=DARK_GREEN)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

takeaways = [
    ("01", "UX ≠ UI", "UX is NOT just about the user interface — it's about the\nentire end-to-end experience across all touchpoints."),
    ("02", "Process First", "Business process redesign is crucial — a prettier UI\nalone won't fix broken workflows."),
    ("03", "Integrate", "Integration across apps eliminates the pain of\ncontext-switching between multiple systems."),
    ("04", "Measure It", "Use a metric-driven approach — measure productivity,\ncost savings, user satisfaction to prove ROI."),
    ("05", "Start Anywhere", "Use a phased implementation — start where you can\nbased on maturity, then expand systematically."),
    ("06", "Partner Up", "UX professionals must closely partner with both\nbusiness teams and technical teams."),
    ("07", "Fewer Vendors", "Long-term: fewer vendor platforms = better unified\nexperience across the organization."),
]

for i, (num, title, desc) in enumerate(takeaways):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + Inches(col * 6.3)
    y = Inches(1.4) + Inches(row * 1.35)
    
    if i == 6:  # Last one centered
        x = Inches(3.5)
    
    # Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.8), Inches(1.15))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = ACCENT_GREEN
    card.line.width = Pt(1.5)
    
    # Number
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.1), y + Inches(0.3), Inches(0.45), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = DARK_GREEN
    badge.line.fill.background()
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Title
    add_text_box(slide, x + Inches(0.65), y + Inches(0.05), Inches(4.8), Inches(0.35),
                 title, font_size=16, bold=True, color=DARK_GREEN)
    # Description
    add_text_box(slide, x + Inches(0.65), y + Inches(0.4), Inches(4.8), Inches(0.65),
                 desc, font_size=11, color=DARK_GRAY)


# ============================================================
# SLIDE 11: EXAM-READY QUICK REFERENCE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)
add_green_bar(slide)
add_bottom_bar(slide)
add_slide_number(slide, 11)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "Exam-Ready Quick Reference", font_size=32, bold=True, color=DARK_GREEN)

exam_tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.3), Inches(2.3), Inches(0.5))
exam_tag.fill.solid()
exam_tag.fill.fore_color.rgb = RED_ACCENT
exam_tag.line.fill.background()
tf = exam_tag.text_frame
p = tf.paragraphs[0]
p.text = "📝 MEMORIZE THIS"
p.font.size = Pt(14)
p.font.color.rgb = WHITE
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(4), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# Box 1: UE Definition
box1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.8), Inches(1.5))
box1.fill.solid()
box1.fill.fore_color.rgb = RGBColor(234, 250, 234)
box1.line.color.rgb = DARK_GREEN
box1.line.width = Pt(2)

add_text_box(slide, Inches(1.0), Inches(1.4), Inches(5.4), Inches(0.3),
             "Definition — Unified Experience =", font_size=14, bold=True, color=DARK_GREEN)
add_text_box(slide, Inches(1.0), Inches(1.7), Inches(5.4), Inches(0.8),
             "Consistency + Efficiency + Personalization + Integration + Optimization\n\nA comprehensive approach to delivering end-to-end experience across multiple applications in a business context.",
             font_size=12, color=DARK_GRAY)

# Box 2: Framework
box2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.3), Inches(5.8), Inches(1.5))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(214, 234, 248)
box2.line.color.rgb = BLUE_ACCENT
box2.line.width = Pt(2)

add_text_box(slide, Inches(7.2), Inches(1.4), Inches(5.4), Inches(0.3),
             "Implementation Framework — 2 Vectors:", font_size=14, bold=True, color=BLUE_ACCENT)
add_text_box(slide, Inches(7.2), Inches(1.7), Inches(5.4), Inches(0.8),
             "1. Organization (Top-Down):\n   → UX Culture + Organizational Commitment\n2. Design (Bottom-Up):\n   → UX Standards, UI Assets, Integration Framework,\n      UX Practices, Vendor Strategy, Governance",
             font_size=11, color=DARK_GRAY)

# Box 3: Klocek's Model
box3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.0), Inches(5.8), Inches(2.0))
box3.fill.solid()
box3.fill.fore_color.rgb = RGBColor(255, 248, 240)
box3.line.color.rgb = RGBColor(230, 126, 34)
box3.line.width = Pt(2)

add_text_box(slide, Inches(1.0), Inches(3.1), Inches(5.4), Inches(0.3),
             "Klocek's 5-Level Pyramid (Previous Model):", font_size=14, bold=True, color=RGBColor(230, 126, 34))
add_text_box(slide, Inches(1.0), Inches(3.4), Inches(5.4), Inches(1.4),
             "Level 1: Consistent Visual Design\nLevel 2: Consistent Interaction Behavior\nLevel 3: Rework Products for User Needs\nLevel 4: Design Unified Experience\nLevel 5: Transform Organization\n\nWeakness: Too linear, overemphasizes visual design, doesn't work well with vendor solutions.",
             font_size=11, color=DARK_GRAY)

# Box 4: Case Study Numbers
box4 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(3.0), Inches(5.8), Inches(2.0))
box4.fill.solid()
box4.fill.fore_color.rgb = RGBColor(255, 235, 235)
box4.line.color.rgb = RED_ACCENT
box4.line.width = Pt(2)

add_text_box(slide, Inches(7.2), Inches(3.1), Inches(5.4), Inches(0.3),
             "Case Study Key Numbers (REMEMBER!):", font_size=14, bold=True, color=RED_ACCENT)
add_text_box(slide, Inches(7.2), Inches(3.4), Inches(5.4), Inches(1.4),
             "• Manual Steps: 35 → 12\n• Applications: 6 → 1\n• Clicks per Transaction: 69 → 17\n• Task Time: 10 min → 5 min\n• Cost Savings per Buyer: $0 → $1.2 Million\n• Training Time: Dozens of hours → 2 hours\n• System Trust: <10% → 83%",
             font_size=12, color=DARK_GRAY)

# Box 5: Design Principles
box5 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.2), Inches(5.8), Inches(1.7))
box5.fill.solid()
box5.fill.fore_color.rgb = VERY_LIGHT_GRAY
box5.line.color.rgb = DARK_GREEN
box5.line.width = Pt(2)

add_text_box(slide, Inches(1.0), Inches(5.3), Inches(5.4), Inches(0.3),
             "6 Design Principles:", font_size=14, bold=True, color=DARK_GREEN)
add_text_box(slide, Inches(1.0), Inches(5.6), Inches(5.4), Inches(1.2),
             "1. Standardize the process\n2. Integrate Business Intelligence\n3. Reduce siloed systems\n4. Collaborate online\n5. Eliminate rework\n6. Allow self-service for customers",
             font_size=12, color=DARK_GRAY)

# Box 6: Norman & Nielsen
box6 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(5.2), Inches(5.8), Inches(1.7))
box6.fill.solid()
box6.fill.fore_color.rgb = RGBColor(240, 234, 248)
box6.line.color.rgb = RGBColor(155, 89, 182)
box6.line.width = Pt(2)

add_text_box(slide, Inches(7.2), Inches(5.3), Inches(5.4), Inches(0.3),
             "Key Definitions & Tech:", font_size=14, bold=True, color=RGBColor(155, 89, 182))
add_text_box(slide, Inches(7.2), Inches(5.6), Inches(5.4), Inches(1.2),
             'Norman & Nielsen: "UX encompasses ALL aspects of the end-user\'s interaction with the company, its services, and its products."\n\nTech Stack: BPMS (Business Process Management Suite) + SOA (Service Oriented Architecture)\n\nNew System: BOMA (Bill of Materials Application)',
             font_size=11, color=DARK_GRAY)


# ============================================================
# SLIDE 12: THANK YOU
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg_white(slide)

# Full green top section
top_section = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(3.5))
top_section.fill.solid()
top_section.fill.fore_color.rgb = DARK_GREEN
top_section.line.fill.background()

add_text_box(slide, Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.0),
             "Thank You!", font_size=54, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER, font_name='Calibri')

add_text_box(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
             "Questions & Discussion", font_size=24, color=RGBColor(144, 238, 144), alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(1.0), Inches(2.7), Inches(11.3), Inches(0.4),
             "Designing for a Unified Experience: A New Perspective and a Case Study — Xu & Furie, Intel Corporation",
             font_size=13, color=RGBColor(200, 230, 200), alignment=PP_ALIGN.CENTER)

# Bottom section with group info
add_text_box(slide, Inches(1.5), Inches(4.0), Inches(10.3), Inches(0.5),
             "Presented by:", font_size=20, bold=True, color=DARK_GREEN, alignment=PP_ALIGN.CENTER)

students_end = [
    "[Student Name 1] — Roll: [XXXXXXX]",
    "[Student Name 2] — Roll: [XXXXXXX]",
    "[Student Name 3] — Roll: [XXXXXXX]",
    "[Student Name 4] — Roll: [XXXXXXX]",
]

for i, s in enumerate(students_end):
    add_text_box(slide, Inches(4.0), Inches(4.6) + Inches(i * 0.4), Inches(5.3), Inches(0.35),
                 s, font_size=15, color=DARK_GRAY, alignment=PP_ALIGN.CENTER)

# Course info at bottom
bottom_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.8), prs.slide_width, Inches(0.7))
bottom_bar.fill.solid()
bottom_bar.fill.fore_color.rgb = DARK_GREEN
bottom_bar.line.fill.background()

add_text_box(slide, Inches(0.5), Inches(6.85), Inches(12.3), Inches(0.5),
             "CSE [Course No.] — [Course Name]  |  [Teacher Name]  |  KUET  |  February 2026",
             font_size=13, color=WHITE, alignment=PP_ALIGN.CENTER)


# ============================================================
# SAVE
# ============================================================
output_path = r'd:\KUET CSE-2K21\3-2\mobile_presentation\Unified_Experience_Design_Presentation.pptx'
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
